"""Report builders: self-contained HTML (print-to-PDF) and DOCX downloads.

Used for the teacher's full survey report and the participant's own transcript.
Generated images stored in ``media_assets`` are embedded inline so reports stay
self-contained after provider URLs expire.
"""

import base64
import html
import io
import json
import re
from datetime import datetime, timezone
from typing import Callable, Optional

import httpx

ASSET_RE = re.compile(r"^/api/assets/([0-9a-fA-F-]{36})$")


def _esc(s) -> str:
    return html.escape(str(s if s is not None else ""))


def parse_tool_events(content: str) -> list:
    try:
        return json.loads(content[len("[TOOL_EVENTS]"):])
    except Exception:
        return []


def transcript_items(messages: list) -> list:
    """Normalise stored messages into renderable items (text, media, buttons)."""
    items = []
    for m in messages:
        content = m.content if hasattr(m, "content") else m["content"]
        role = m.role if hasattr(m, "role") else m["role"]
        created = m.created_at if hasattr(m, "created_at") else m.get("created_at")
        if content.startswith("[TOOL_EVENTS]"):
            for ev in parse_tool_events(content):
                if ev.get("t") == "media":
                    items.append({"kind": "media", "role": "assistant", "event": ev, "created_at": created})
                elif ev.get("t") == "buttons":
                    items.append({"kind": "buttons", "role": "assistant", "event": ev, "created_at": created})
                elif ev.get("t") == "interactive":
                    items.append({"kind": "interactive", "role": "assistant", "event": ev, "created_at": created})
            continue
        if content.startswith("(The participant has just joined"):
            continue
        if content == "(presented interactive content)":
            continue
        items.append({"kind": "text", "role": role, "text": content, "created_at": created})
    return items


def _image_data_uri(url: str, asset_loader: Optional[Callable], fetch_external: bool) -> Optional[str]:
    """Return a data: URI for an image url, using the DB for local assets."""
    if not url:
        return None
    m = ASSET_RE.match(url)
    if m and asset_loader:
        asset = asset_loader(m.group(1))
        if asset:
            return f"data:{asset.mime_type};base64,{base64.b64encode(asset.data).decode()}"
        return None
    if fetch_external and url.startswith("http"):
        try:
            resp = httpx.get(url, timeout=8.0, follow_redirects=True)
            if resp.status_code == 200 and resp.headers.get("content-type", "").startswith("image/"):
                mime = resp.headers["content-type"].split(";")[0]
                return f"data:{mime};base64,{base64.b64encode(resp.content).decode()}"
        except Exception:
            return None
    return url


# ─────────────────────────── HTML ───────────────────────────

REPORT_CSS = """
*{box-sizing:border-box}
body{font-family:'Segoe UI',-apple-system,BlinkMacSystemFont,Roboto,Helvetica,Arial,sans-serif;color:#1f2937;margin:0;background:#f4f6fa}
.page{max-width:860px;margin:0 auto;padding:40px 36px;background:#fff}
h1{font-size:1.7rem;margin:0 0 4px;color:#0f172a}
h2{font-size:1.15rem;margin:32px 0 10px;color:#0f172a;border-bottom:2px solid #e2e8f0;padding-bottom:6px}
h3{font-size:1rem;margin:20px 0 8px;color:#1e293b}
.sub{color:#64748b;font-size:0.9rem;margin-bottom:24px}
.meta{display:flex;flex-wrap:wrap;gap:10px;margin:14px 0 6px}
.stat{flex:1;min-width:120px;background:#f8fafc;border:1px solid #e2e8f0;border-radius:10px;padding:12px 14px}
.stat .n{font-size:1.4rem;font-weight:700;color:#1d4ed8}
.stat .l{font-size:0.75rem;color:#64748b;text-transform:uppercase;letter-spacing:.04em}
table{border-collapse:collapse;width:100%;font-size:0.86rem;margin:8px 0}
th,td{border:1px solid #e2e8f0;padding:7px 10px;text-align:left;vertical-align:top}
th{background:#f1f5f9;font-weight:600}
.bar{height:10px;background:#e2e8f0;border-radius:6px;overflow:hidden;margin-top:4px}
.bar span{display:block;height:100%}
.pos{background:#16a34a}.neu{background:#d97706}.neg{background:#dc2626}
.msg{margin:8px 0;padding:10px 14px;border-radius:12px;max-width:88%;font-size:0.92rem;line-height:1.5;white-space:pre-wrap;word-wrap:break-word}
.msg.assistant{background:#f1f5f9;border:1px solid #e2e8f0;margin-right:auto}
.msg.user{background:#1d4ed8;color:#fff;margin-left:auto}
.msg .who{font-size:0.7rem;text-transform:uppercase;letter-spacing:.05em;opacity:.7;margin-bottom:3px;white-space:normal}
.media{margin:10px 0;max-width:88%}
.media img{max-width:100%;border-radius:10px;border:1px solid #e2e8f0;display:block}
.media .cap{font-size:0.8rem;color:#64748b;margin-top:4px}
.opts{display:flex;flex-wrap:wrap;gap:6px;margin:6px 0 10px}
.opts span{border:1px solid #cbd5e1;border-radius:99px;padding:3px 10px;font-size:0.8rem;color:#334155}
.participant{page-break-before:always;padding-top:8px}
.participant:first-of-type{page-break-before:auto}
.brief{background:#f8fafc;border-left:4px solid #1d4ed8;padding:12px 16px;border-radius:0 8px 8px 0;white-space:pre-wrap;font-size:0.92rem}
.toolbar{position:sticky;top:0;background:#0f172a;color:#fff;padding:10px 16px;display:flex;gap:10px;align-items:center;font-size:0.88rem}
.toolbar button{background:#2563eb;border:none;color:#fff;padding:8px 14px;border-radius:8px;cursor:pointer;font-weight:600}
.toolbar a{color:#cbd5e1;text-decoration:none;margin-left:auto}
.small{font-size:0.8rem;color:#64748b}
@media print{.toolbar{display:none}.page{padding:0;max-width:none}body{background:#fff}}
"""


INTERACTIVE_LABELS = {"mcq": "Multiple choice", "fill_blank": "Fill in the blanks", "order": "Put in order",
                      "match": "Match the pairs", "scale": "Scale"}


def interactive_summary(ev: dict) -> str:
    """One-line description of an interactive check and its answer key, for reports."""
    kind = ev.get("kind", "")
    label = INTERACTIVE_LABELS.get(kind, "Interactive")
    q = ev.get("question") or ""
    detail = ""
    if kind == "mcq":
        opts = ev.get("options") or []
        c = ev.get("correct")
        detail = " / ".join(f"{'✓ ' if c == i else ''}{o}" for i, o in enumerate(opts))
    elif kind == "fill_blank":
        detail = f"{ev.get('text', '')} — answers: {', '.join(ev.get('answers') or [])}"
    elif kind == "order":
        detail = "correct order: " + " → ".join(ev.get("items") or [])
    elif kind == "match":
        detail = "; ".join(f"{p.get('left')} ↔ {p.get('right')}" for p in ev.get("pairs") or [])
    elif kind == "scale":
        detail = f"{ev.get('min_label', '')} … {ev.get('max_label', '')}"
    return f"{label}: {q}" + (f" [{detail}]" if detail else "")


def _render_items(items: list, asset_loader, fetch_external: bool) -> str:
    out = []
    for it in items:
        if it["kind"] == "text":
            who = "Participant" if it["role"] == "user" else "Facilitator"
            out.append(f'<div class="msg {it["role"]}"><div class="who">{who}</div>{_esc(it["text"])}</div>')
        elif it["kind"] == "media":
            ev = it["event"]
            if ev.get("type") == "image":
                src = _image_data_uri(ev.get("url", ""), asset_loader, fetch_external)
                if src:
                    cap = ev.get("caption") or ev.get("alt") or ""
                    out.append(f'<div class="media"><img src="{_esc(src)}" alt="{_esc(cap)}">'
                               f'{f"<div class=cap>{_esc(cap)}</div>" if cap else ""}</div>')
            elif ev.get("type") == "video":
                out.append(f'<div class="small">[Video shown: {_esc(ev.get("caption") or ev.get("url", ""))}]</div>')
        elif it["kind"] == "buttons":
            ev = it["event"]
            labels = "".join(f"<span>{_esc(o.get('label', ''))}</span>" for o in ev.get("options", []))
            q = ev.get("question") or ""
            out.append(f'<div class="small">Options offered{": " + _esc(q) if q else ""}</div><div class="opts">{labels}</div>')
        elif it["kind"] == "interactive":
            out.append(f'<div class="small">Interactive check — {_esc(interactive_summary(it["event"]))}</div>')
    return "\n".join(out)


def build_survey_report_html(survey, participants: list, stats: dict, insights: Optional[dict],
                             asset_loader=None, fetch_external: bool = False) -> str:
    now = datetime.now(timezone.utc).strftime("%d %b %Y, %H:%M UTC")
    parts = [f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{_esc(survey.title)} – Report</title>"
             f"<style>{REPORT_CSS}</style></head><body>"
             f"<div class='toolbar'><button onclick='window.print()'>Print / Save as PDF</button>"
             f"<span>Use your browser's print dialog and choose “Save as PDF”.</span>"
             f"<a href='javascript:window.close()'>Close</a></div><div class='page'>"]
    parts.append(f"<h1>{_esc(survey.title)}</h1>")
    parts.append(f"<div class='sub'>{_esc(survey.topic)}<br>Survey code <b>{_esc(survey.survey_code)}</b> · "
                 f"Status {_esc(survey.status.value)} · Report generated {now}</div>")

    parts.append("<div class='meta'>")
    for key, label in [("total_participants", "Participants"), ("completed_participants", "Completed"),
                       ("active_participants", "In progress")]:
        parts.append(f"<div class='stat'><div class='n'>{stats.get(key, 0)}</div><div class='l'>{label}</div></div>")
    avg = stats.get("avg_completion_seconds") or 0
    parts.append(f"<div class='stat'><div class='n'>{round(avg / 60, 1) if avg else '—'}</div><div class='l'>Avg minutes</div></div>")
    parts.append("</div>")

    if survey.briefing_text or survey.briefing_url:
        parts.append("<h2>Task briefing</h2>")
        if survey.briefing_title:
            parts.append(f"<h3>{_esc(survey.briefing_title)}</h3>")
        if survey.briefing_text:
            parts.append(f"<div class='brief'>{_esc(survey.briefing_text)}</div>")
        if survey.briefing_url:
            parts.append(f"<p class='small'>Briefing material: {_esc(survey.briefing_url)}</p>")

    if survey.questions:
        parts.append("<h2>Questions asked</h2>")
        parts.append(f"<div class='brief'>{_esc(survey.questions)}</div>")

    if insights and not insights.get("error"):
        sent = insights.get("sentiment") or {}
        total = sum(int(sent.get(k, 0) or 0) for k in ("positive", "neutral", "negative")) or 1
        parts.append("<h2>AI insights</h2><h3>Sentiment</h3><table><tr><th>Sentiment</th><th>Count</th><th style='width:45%'>Share</th></tr>")
        for k, cls in (("positive", "pos"), ("neutral", "neu"), ("negative", "neg")):
            v = int(sent.get(k, 0) or 0)
            pct = round(v * 100 / total)
            parts.append(f"<tr><td>{k.title()}</td><td>{v}</td><td>{pct}%<div class='bar'><span class='{cls}' style='width:{pct}%'></span></div></td></tr>")
        parts.append("</table>")
        themes = insights.get("themes") or []
        if themes:
            parts.append("<h3>Top themes</h3><table><tr><th>#</th><th>Theme</th><th>Mentions</th></tr>")
            for i, t in enumerate(themes[:10], 1):
                parts.append(f"<tr><td>{i}</td><td>{_esc(t.get('name'))}</td><td>{_esc(t.get('count'))}</td></tr>")
            parts.append("</table>")
        pins = insights.get("participants") or []
        if pins:
            parts.append("<h3>Per-participant summary</h3><table><tr><th>Participant</th><th>Sentiment</th><th>Engagement</th><th>Themes</th></tr>")
            for p in pins:
                parts.append(f"<tr><td>{_esc(p.get('id'))}</td><td>{_esc(p.get('sentiment'))}</td>"
                             f"<td>{_esc(p.get('engagement'))}/10</td><td>{_esc(', '.join(p.get('themes') or []))}</td></tr>")
            parts.append("</table>")

    parts.append("<h2>Transcripts</h2>")
    if not participants:
        parts.append("<p class='small'>No participants yet.</p>")
    for p in participants:
        items = transcript_items(p["messages"])
        if not items:
            continue
        label = p.get("contact_name") or f"Participant {p['id'][:8]}"
        dur = f"{round(p['duration_seconds'] / 60, 1)} min" if p.get("duration_seconds") else "in progress"
        started = p.get("started_at", "")
        parts.append(f"<div class='participant'><h3>{_esc(label)}</h3>"
                     f"<div class='small'>ID {p['id'][:8]} · {_esc(p['status'])} · {dur} · started {_esc(started[:16].replace('T', ' '))}"
                     f"{' · ' + _esc(p['contact_email']) if p.get('contact_email') else ''}</div>")
        parts.append(_render_items(items, asset_loader, fetch_external))
        parts.append("</div>")

    parts.append("</div></body></html>")
    return "".join(parts)


def build_participant_report_html(survey, participant, messages: list, asset_loader=None) -> str:
    items = transcript_items(messages)
    now = datetime.now(timezone.utc).strftime("%d %b %Y")
    label = participant.contact_name or "Participant"
    parts = [f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{_esc(survey.title)} – My responses</title>"
             f"<style>{REPORT_CSS}</style></head><body>"
             f"<div class='toolbar'><button onclick='window.print()'>Print / Save as PDF</button>"
             f"<span>Choose “Save as PDF” in the print dialog to keep a copy.</span>"
             f"<a href='javascript:window.close()'>Close</a></div><div class='page'>"]
    parts.append(f"<h1>{_esc(survey.title)}</h1><div class='sub'>{_esc(survey.topic)}<br>"
                 f"My responses · {_esc(label)} · {now}</div>")
    if survey.briefing_text:
        parts.append(f"<h2>Task</h2><div class='brief'>{_esc(survey.briefing_text)}</div>")
    parts.append("<h2>Conversation</h2>")
    parts.append(_render_items(items, asset_loader, fetch_external=False))
    parts.append("</div></body></html>")
    return "".join(parts)


# ─────────────────────────── DOCX ───────────────────────────

def build_survey_report_docx(survey, participants: list, stats: dict, insights: Optional[dict],
                             asset_loader=None, fetch_external: bool = False) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)

    doc.add_heading(survey.title, level=0)
    p = doc.add_paragraph(survey.topic)
    p.runs[0].italic = True
    doc.add_paragraph(
        f"Survey code: {survey.survey_code}   ·   Status: {survey.status.value}   ·   "
        f"Generated: {datetime.now(timezone.utc).strftime('%d %b %Y, %H:%M UTC')}"
    )

    doc.add_heading("Summary", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, t in enumerate(["Participants", "Completed", "In progress", "Avg minutes"]):
        hdr[i].text = t
    avg = stats.get("avg_completion_seconds") or 0
    row = table.add_row().cells
    row[0].text = str(stats.get("total_participants", 0))
    row[1].text = str(stats.get("completed_participants", 0))
    row[2].text = str(stats.get("active_participants", 0))
    row[3].text = str(round(avg / 60, 1)) if avg else "—"

    if survey.briefing_text or survey.briefing_url:
        doc.add_heading("Task briefing", level=1)
        if survey.briefing_title:
            doc.add_paragraph(survey.briefing_title).runs[0].bold = True
        if survey.briefing_text:
            doc.add_paragraph(survey.briefing_text)
        if survey.briefing_url:
            doc.add_paragraph(f"Briefing material: {survey.briefing_url}")

    if survey.questions:
        doc.add_heading("Questions asked", level=1)
        doc.add_paragraph(survey.questions)

    if insights and not insights.get("error"):
        doc.add_heading("AI insights", level=1)
        sent = insights.get("sentiment") or {}
        total = sum(int(sent.get(k, 0) or 0) for k in ("positive", "neutral", "negative")) or 1
        t = doc.add_table(rows=1, cols=3)
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(["Sentiment", "Count", "Share"]):
            t.rows[0].cells[i].text = h
        for k in ("positive", "neutral", "negative"):
            v = int(sent.get(k, 0) or 0)
            r = t.add_row().cells
            r[0].text, r[1].text, r[2].text = k.title(), str(v), f"{round(v * 100 / total)}%"
        themes = insights.get("themes") or []
        if themes:
            doc.add_heading("Top themes", level=2)
            for i, th in enumerate(themes[:10], 1):
                doc.add_paragraph(f"{th.get('name')} — {th.get('count')} mentions", style="List Number")
        pins = insights.get("participants") or []
        if pins:
            doc.add_heading("Per-participant summary", level=2)
            t = doc.add_table(rows=1, cols=4)
            t.style = "Light Grid Accent 1"
            for i, h in enumerate(["Participant", "Sentiment", "Engagement", "Themes"]):
                t.rows[0].cells[i].text = h
            for pi in pins:
                r = t.add_row().cells
                r[0].text = str(pi.get("id", ""))
                r[1].text = str(pi.get("sentiment", ""))
                r[2].text = f"{pi.get('engagement', '')}/10"
                r[3].text = ", ".join(pi.get("themes") or [])

    doc.add_heading("Transcripts", level=1)
    first = True
    for p in participants:
        items = transcript_items(p["messages"])
        if not items:
            continue
        if not first:
            doc.add_page_break()
        first = False
        label = p.get("contact_name") or f"Participant {p['id'][:8]}"
        doc.add_heading(label, level=2)
        dur = f"{round(p['duration_seconds'] / 60, 1)} min" if p.get("duration_seconds") else "in progress"
        meta = f"ID {p['id'][:8]} · {p['status']} · {dur}"
        if p.get("contact_email"):
            meta += f" · {p['contact_email']}"
        mp = doc.add_paragraph(meta)
        mp.runs[0].font.size = Pt(8.5)
        mp.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        for it in items:
            if it["kind"] == "text":
                who = "Participant" if it["role"] == "user" else "Facilitator"
                para = doc.add_paragraph()
                run = para.add_run(f"{who}: ")
                run.bold = True
                run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8) if it["role"] == "user" else RGBColor(0x33, 0x41, 0x55)
                para.add_run(it["text"])
            elif it["kind"] == "media" and it["event"].get("type") == "image":
                ev = it["event"]
                uri = _image_data_uri(ev.get("url", ""), asset_loader, fetch_external)
                if uri and uri.startswith("data:"):
                    try:
                        raw = base64.b64decode(uri.split(",", 1)[1])
                        doc.add_picture(io.BytesIO(raw), width=Inches(4.5))
                        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
                    except Exception:
                        doc.add_paragraph("[image could not be embedded]")
                cap = ev.get("caption") or ev.get("alt")
                if cap:
                    cp = doc.add_paragraph(cap)
                    cp.runs[0].italic = True
                    cp.runs[0].font.size = Pt(8.5)
            elif it["kind"] == "buttons":
                labels = ", ".join(o.get("label", "") for o in it["event"].get("options", []))
                bp = doc.add_paragraph(f"[Options offered: {labels}]")
                bp.runs[0].font.size = Pt(8.5)
                bp.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            elif it["kind"] == "interactive":
                ip = doc.add_paragraph(f"[Interactive check — {interactive_summary(it['event'])}]")
                ip.runs[0].font.size = Pt(8.5)
                ip.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
